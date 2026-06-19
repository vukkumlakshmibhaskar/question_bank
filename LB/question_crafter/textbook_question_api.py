import io
import json
import os
import re
import shutil
import uuid
from pathlib import Path

import pandas as pd
from fastapi import Body, FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font

from .question_engine import (
    AIUnavailableError,
    EXPORT_COLUMNS,
    GenerationSettings,
    LessonBlock,
    combine_texts,
    clean_export_text,
    clean_question_export_text,
    detect_lessons,
    generate_answers_for_rows,
    generate_visual_questions_from_images,
    generate_questions,
    generate_lesson_questions,
    keywords,
)


BASE_DIR = Path(__file__).resolve().parents[1]
EXPORT_DIR = BASE_DIR / "workspace" / "question_exports"
SESSION_DIR = BASE_DIR / "workspace" / "review_sessions"
TEXTBOOK_DIR = BASE_DIR / "workspace" / "textbook_sessions"
PAGE_IMAGE_DIR = TEXTBOOK_DIR / "pages"
EXPORT_DIR.mkdir(parents=True, exist_ok=True)
SESSION_DIR.mkdir(parents=True, exist_ok=True)
TEXTBOOK_DIR.mkdir(parents=True, exist_ok=True)
PAGE_IMAGE_DIR.mkdir(parents=True, exist_ok=True)
CROP_IMAGE_DIR = TEXTBOOK_DIR / "crops"
CROP_IMAGE_DIR.mkdir(parents=True, exist_ok=True)
ANSWER_REVIEW_MARKER = "Needs answer review"
BLANK_ANSWER_VALUES = {"", "none", "null", "nan", "na", "n/a", "-", "--"}

raw_cors_origins = os.getenv(
    "CORS_ORIGINS",
    "http://localhost:5173,http://127.0.0.1:5173,https://qbparser1.imove4m.com",
)
cors_origins = [origin.strip().rstrip("/") for origin in raw_cors_origins.split(",") if origin.strip()]

app = FastAPI(title="QC Textbook Question Generator", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=cors_origins if cors_origins != ["*"] else ["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)
app.mount("/textbook-pages", StaticFiles(directory=PAGE_IMAGE_DIR), name="textbook-pages")
app.mount("/textbook-crops", StaticFiles(directory=CROP_IMAGE_DIR), name="textbook-crops")

FINAL_EXPORT_COLUMNS = [
    "question_number",
    "question_type",
    "question_image",
    "question_image_file",
    "question",
    "AI answer",
    "bloom_tag",
    "difficulty",
    "lesson_no",
    "lesson_name",
    "subject_name",
]


def safe_path_part(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9_.-]+", "_", value or "file").strip("._")
    return cleaned[:80] or "file"


def render_pdf_page_images(data: bytes, session_id: str, filename: str) -> list[str]:
    try:
        import fitz
    except Exception:
        return []
    safe_name = safe_path_part(Path(filename or "textbook").stem)
    target_dir = PAGE_IMAGE_DIR / session_id / safe_name
    target_dir.mkdir(parents=True, exist_ok=True)
    urls: list[str] = []
    try:
        doc = fitz.open(stream=data, filetype="pdf")
        matrix = fitz.Matrix(1.6, 1.6)
        for page_index, page in enumerate(doc, start=1):
            image_path = target_dir / f"page_{page_index:03d}.png"
            if not image_path.exists():
                pix = page.get_pixmap(matrix=matrix, alpha=False)
                pix.save(image_path)
            urls.append(f"/textbook-pages/{session_id}/{safe_name}/page_{page_index:03d}.png")
    except Exception:
        return []
    return urls


def auto_crop_page_diagrams(page_image_path: Path, session_id: str, page_url: str) -> list[str]:
    try:
        import cv2
        import numpy as np
    except Exception:
        return []

    image = cv2.imread(str(page_image_path))
    if image is None:
        return []

    height, width = image.shape[:2]
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    blurred = cv2.GaussianBlur(gray, (3, 3), 0)
    edges = cv2.Canny(blurred, 60, 160)
    kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (9, 9))
    closed = cv2.morphologyEx(edges, cv2.MORPH_CLOSE, kernel, iterations=2)
    contours, _ = cv2.findContours(closed, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)

    candidates = []
    page_area = width * height
    for contour in contours:
        x, y, w, h = cv2.boundingRect(contour)
        area = w * h
        if area < page_area * 0.015 or area > page_area * 0.55:
            continue
        if w < width * 0.12 or h < height * 0.06:
            continue
        if y < height * 0.05 or y + h > height * 0.96:
            continue
        aspect = w / max(h, 1)
        if aspect > 7 or aspect < 0.12:
            continue
        roi = gray[y:y + h, x:x + w]
        dark_ratio = float(np.mean(roi < 210))
        if dark_ratio < 0.015:
            continue
        # Prefer boxed/visual regions over text-only paragraphs.
        score = area * (1 + min(dark_ratio, 0.35))
        candidates.append((score, x, y, w, h))

    candidates.sort(reverse=True)
    kept = []
    for _, x, y, w, h in candidates:
        box = (x, y, x + w, y + h)
        if any(_iou(box, existing) > 0.25 for existing in kept):
            continue
        kept.append(box)
        if len(kept) >= 4:
            break

    crop_urls = []
    target_dir = CROP_IMAGE_DIR / safe_path_part(session_id) / "auto"
    target_dir.mkdir(parents=True, exist_ok=True)
    page_stem = safe_path_part(Path(page_url).stem)
    for index, (left, top, right, bottom) in enumerate(kept, start=1):
        pad_x = int(width * 0.015)
        pad_y = int(height * 0.015)
        left = max(0, left - pad_x)
        top = max(0, top - pad_y)
        right = min(width, right + pad_x)
        bottom = min(height, bottom + pad_y)
        crop = image[top:bottom, left:right]
        filename = f"{page_stem}_auto_{index}.png"
        cv2.imwrite(str(target_dir / filename), crop)
        crop_urls.append(f"/textbook-crops/{safe_path_part(session_id)}/auto/{filename}")
    return crop_urls


def _iou(a: tuple[int, int, int, int], b: tuple[int, int, int, int]) -> float:
    ax1, ay1, ax2, ay2 = a
    bx1, by1, bx2, by2 = b
    ix1, iy1 = max(ax1, bx1), max(ay1, by1)
    ix2, iy2 = min(ax2, bx2), min(ay2, by2)
    if ix2 <= ix1 or iy2 <= iy1:
        return 0.0
    intersection = (ix2 - ix1) * (iy2 - iy1)
    area_a = (ax2 - ax1) * (ay2 - ay1)
    area_b = (bx2 - bx1) * (by2 - by1)
    return intersection / max(area_a + area_b - intersection, 1)


def build_auto_crop_index(session_id: str) -> dict[str, list[str]]:
    crop_index: dict[str, list[str]] = {}
    root = PAGE_IMAGE_DIR / safe_path_part(session_id)
    if not root.exists():
        return crop_index
    for image_path in sorted(root.rglob("page_*.png")):
        relative = image_path.relative_to(PAGE_IMAGE_DIR).as_posix()
        page_url = f"/textbook-pages/{relative}"
        crops = auto_crop_page_diagrams(image_path, session_id, page_url)
        if crops:
            crop_index[page_url] = crops
    return crop_index


def flatten_pdf_outline(reader) -> list[tuple[str, int]]:
    items: list[tuple[str, int]] = []

    def walk(outline_items):
        for item in outline_items:
            if isinstance(item, list):
                walk(item)
                continue
            title = str(getattr(item, "title", item)).strip()
            try:
                page_index = reader.get_destination_page_number(item)
            except Exception:
                continue
            items.append((title, page_index))

    try:
        walk(reader.outline)
    except Exception:
        return []
    return items


def infer_lesson_title(page_text: str, lesson_number: int, fallback: str) -> str:
    lines = [line.strip() for line in (page_text or "").splitlines() if line.strip()]
    ignored = {"notes", "physics"}
    for index, line in enumerate(lines[:20]):
        if line == str(lesson_number):
            title_lines = []
            for candidate in lines[index + 1:index + 5]:
                compact = candidate.lower()
                if compact in ignored or compact.startswith("module"):
                    if title_lines:
                        break
                    continue
                if len(candidate.split()) > 8 or re.search(r"[.!?]$", candidate):
                    break
                title_lines.append(candidate)
                if len(title_lines) >= 2:
                    break
            if title_lines:
                return " ".join(title_lines).replace("  ", " ")
    return fallback


def page_marker(page_index: int, page_image_urls: list[str] | None = None) -> str:
    image_url = page_image_urls[page_index] if page_image_urls and page_index < len(page_image_urls) else ""
    suffix = f" IMAGE:{image_url}" if image_url else ""
    return f"[[PAGE:{page_index + 1}{suffix}]]"


def extract_pdf_text(reader, page_image_urls: list[str] | None = None) -> str:
    lesson_marks = []
    for title, page_index in flatten_pdf_outline(reader):
        match = re.search(r"\blesson\s*[-:]?\s*(\d{1,3})\b", title, flags=re.IGNORECASE)
        if match:
            lesson_marks.append((int(match.group(1)), page_index))

    if not lesson_marks:
        return "\n".join(f"{page_marker(index, page_image_urls)}\n{page.extract_text() or ''}" for index, page in enumerate(reader.pages))

    lesson_marks = sorted(dict((num, page) for num, page in lesson_marks).items(), key=lambda item: item[1])
    sections = []
    for index, (lesson_number, start_page) in enumerate(lesson_marks):
        end_page = lesson_marks[index + 1][1] if index + 1 < len(lesson_marks) else len(reader.pages)
        page_texts = [f"{page_marker(p, page_image_urls)}\n{reader.pages[p].extract_text() or ''}" for p in range(start_page, end_page)]
        fallback = f"Lesson{lesson_number:02d}"
        lesson_title = infer_lesson_title(page_texts[0] if page_texts else "", lesson_number, fallback)
        body = "\n".join(page_texts)
        body = re.sub(r"(?im)^\s*((?:lesson|chapter)\s*[-:]?\s*\d{1,3}\b)", r"Reference \1", body)
        sections.append(f"Lesson {lesson_number:02d}\n{lesson_title}\n{body}")
    return "\n\n".join(sections)


async def read_upload(upload: UploadFile, session_id: str | None = None) -> tuple[str, str]:
    name = upload.filename or "textbook"
    suffix = Path(name).suffix.lower()
    data = await upload.read()
    if suffix in {".txt", ".md"}:
        return name, data.decode("utf-8", errors="ignore")
    if suffix == ".pdf":
        from pypdf import PdfReader
        page_image_urls = render_pdf_page_images(data, session_id or "adhoc", name)
        reader = PdfReader(io.BytesIO(data))
        return name, extract_pdf_text(reader, page_image_urls)
    if suffix == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        return name, "\n".join(p.text for p in doc.paragraphs)
    raise ValueError(f"Unsupported file type: {suffix or name}. Use PDF, DOCX, TXT, or MD.")


@app.get("/health")
def health():
    return {"status": "ok", "service": "textbook-question-generator"}


def public_image_to_path(image_url: str) -> Path | None:
    clean_url = str(image_url or "").split("?", 1)[0].strip()
    if clean_url.startswith("/textbook-crops/"):
        relative = clean_url.replace("/textbook-crops/", "", 1).strip("/")
        return (CROP_IMAGE_DIR / relative).resolve()
    if clean_url.startswith("/textbook-pages/"):
        relative = clean_url.replace("/textbook-pages/", "", 1).strip("/")
        return (PAGE_IMAGE_DIR / relative).resolve()
    return None


def cropped_image_filename(image_url: str) -> str:
    clean_url = str(image_url or "").split("?", 1)[0].strip()
    if not clean_url.startswith("/textbook-crops/"):
        return ""
    return Path(clean_url).name


def export_answer_text(value: str) -> str:
    cleaned = clean_export_text(value)
    cleaned = re.sub(r"(?i)<br\s*/?>", "\n", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    if cleaned.lower() in BLANK_ANSWER_VALUES:
        return ANSWER_REVIEW_MARKER
    return cleaned or ANSWER_REVIEW_MARKER


def final_export_rows(workbook_rows: dict[str, list[dict]]) -> dict[str, list[dict]]:
    final_rows: dict[str, list[dict]] = {}
    for sheet_name, rows in workbook_rows.items():
        final_rows[str(sheet_name)] = []
        for row in rows:
            crop_url = str(row.get("source_image_url") or "")
            image_file = cropped_image_filename(crop_url)
            final_rows[str(sheet_name)].append({
                "question_number": row.get("question_number", ""),
                "question_type": row.get("question_type", ""),
                "question_image": "Embedded image" if image_file else "",
                "question_image_file": image_file,
                "question": clean_question_export_text(row.get("question", "")),
                "AI answer": export_answer_text(row.get("AI answer", "")),
                "bloom_tag": row.get("bloom_tag", ""),
                "difficulty": row.get("difficulty", ""),
                "lesson_no": row.get("lesson_no", ""),
                "lesson_name": row.get("lesson_name", ""),
                "subject_name": row.get("subject_name", ""),
                "__image_url": crop_url if image_file else "",
            })
    return final_rows


def write_workbook(path: Path, workbook_rows: dict[str, list[dict]], final_export: bool = False):
    columns = FINAL_EXPORT_COLUMNS if final_export else EXPORT_COLUMNS
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, sheet_rows in workbook_rows.items():
            safe_sheet_name = str(sheet_name or "Questions")[:31]
            pd.DataFrame(sheet_rows, columns=columns).to_excel(writer, sheet_name=safe_sheet_name, index=False)
            worksheet = writer.sheets[safe_sheet_name]
            worksheet.freeze_panes = "A2"
            worksheet.sheet_view.showGridLines = True
            widths = (
                {
                    "A": 16,
                    "B": 18,
                    "C": 28,
                    "D": 28,
                    "E": 70,
                    "F": 55,
                    "G": 18,
                    "H": 14,
                    "I": 14,
                    "J": 32,
                    "K": 24,
                }
                if final_export
                else {
                    "A": 16,
                    "B": 18,
                    "C": 70,
                    "D": 55,
                    "E": 18,
                    "F": 14,
                    "G": 14,
                    "H": 32,
                    "I": 24,
                    "J": 12,
                    "K": 46,
                    "L": 55,
                }
            )
            for col, width in widths.items():
                worksheet.column_dimensions[col].width = width
            for cell in worksheet[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, max_col=len(columns)):
                for cell in row:
                    if cell.value is None:
                        continue
                    cell.value = str(cell.value)
                    cell.data_type = "s"
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                question_cell = row[4] if final_export else row[2]
                answer_cell = row[5] if final_export else row[3]
                if question_cell.value:
                    line_count = str(question_cell.value).count("\n") + 1
                    worksheet.row_dimensions[question_cell.row].height = min(max(42, line_count * 18), 180)
                if answer_cell.value:
                    line_count = str(answer_cell.value).count("\n") + 1
                    worksheet.row_dimensions[answer_cell.row].height = min(max(worksheet.row_dimensions[answer_cell.row].height or 42, line_count * 18), 180)
            if final_export:
                for excel_row, source_row in enumerate(sheet_rows, start=2):
                    image_path = public_image_to_path(source_row.get("__image_url", ""))
                    if not image_path or not image_path.exists():
                        continue
                    try:
                        xl_image = XLImage(str(image_path))
                        scale = min(170 / max(xl_image.width, 1), 120 / max(xl_image.height, 1), 1)
                        xl_image.width = int(xl_image.width * scale)
                        xl_image.height = int(xl_image.height * scale)
                        worksheet.add_image(xl_image, f"C{excel_row}")
                        worksheet.row_dimensions[excel_row].height = max(120, worksheet.row_dimensions[excel_row].height or 42)
                        worksheet[f"C{excel_row}"].value = ""
                    except Exception:
                        worksheet[f"C{excel_row}"].value = "Image file available"


def clean_workbook_rows(workbook_rows: dict[str, list[dict]]) -> dict[str, list[dict]]:
    cleaned: dict[str, list[dict]] = {}
    for sheet_name, rows in workbook_rows.items():
        if not isinstance(rows, list):
            continue
        sheet_rows = []
        for row in rows:
            if not isinstance(row, dict):
                continue
            clean_row = {column: row.get(column, "") for column in EXPORT_COLUMNS}
            clean_row["question"] = clean_question_export_text(clean_row.get("question", ""))
            clean_row["AI answer"] = export_answer_text(clean_row.get("AI answer", ""))
            clean_row["source_excerpt"] = clean_export_text(clean_row.get("source_excerpt", ""))
            sheet_rows.append(clean_row)
        cleaned[str(sheet_name)] = sheet_rows
    return cleaned


DIAGRAM_TERMS_RE = re.compile(
    r"\b(?:figure|fig\.?|diagram|flow\s*chart|flowchart|graph|chart|map|"
    r"labelled|labeled|draw|sketch|axis|axes|plot|table|timeline|cycle)\b",
    flags=re.IGNORECASE,
)

def row_needs_visual(row: dict) -> bool:
    q_type = str(row.get("question_type", "") or "").strip().lower()
    if q_type not in {"diagram", "graph"}:
        return False
    question = str(row.get("question", "") or "")
    return bool(re.search(r"\b(?:given|shown|following|provided|attached)\s+(?:diagram|graph|chart|figure|flow\s*chart|flowchart|map)\b", question, flags=re.IGNORECASE))


def visual_review_reason(row: dict) -> str:
    q_type = str(row.get("question_type", "") or "").strip().lower()
    if q_type not in {"diagram", "graph"}:
        return ""
    question = str(row.get("question", "") or "")
    if re.search(r"\b(?:draw|sketch|illustrate|construct)\b", question, flags=re.IGNORECASE) and not re.search(r"\b(?:given|shown|following|provided|attached)\b", question, flags=re.IGNORECASE):
        return "Draw/sketch question; no source image should be attached."
    if not row_needs_visual(row):
        return "Image-based wording is not explicit. Keep this as text/drawing unless reviewed."
    return ""


def lesson_crop_urls(lesson_text: str, crop_index: dict[str, list[str]]) -> list[str]:
    if not crop_index:
        return []
    urls: list[str] = []
    page_markers = re.findall(r"\[\[PAGE:(\d+)\s+IMAGE:([^\]]+)\]\]", lesson_text or "")
    page_urls = [image_url.strip() for _, image_url in page_markers if image_url.strip()]
    if not page_urls:
        page_numbers = {page for page, _ in re.findall(r"\[\[PAGE:(\d+)(?:\s+IMAGE:([^\]]+))?\]\]", lesson_text or "")}
        for source_url, crops in crop_index.items():
            for page in page_numbers:
                if source_url.endswith(f"page_{int(page):03d}.png"):
                    urls.extend(crops)
    for page_url in page_urls:
        urls.extend(crop_index.get(page_url) or [])
    return list(dict.fromkeys(urls))


def lesson_source_pages(lesson_text: str, crop_index: dict[str, list[str]]) -> list[dict]:
    pages: list[dict] = []
    seen: set[str] = set()
    for page, image_url in re.findall(r"\[\[PAGE:(\d+)\s+IMAGE:([^\]]+)\]\]", lesson_text or ""):
        image_url = image_url.strip()
        if not image_url or image_url in seen:
            continue
        seen.add(image_url)
        pages.append({
            "page": page,
            "image_url": image_url,
            "crops": crop_index.get(image_url) or [],
        })
    return pages


def source_page_texts(lesson_text: str) -> dict[str, str]:
    pages: dict[str, str] = {}
    matches = list(re.finditer(r"\[\[PAGE:(\d+)(?:\s+IMAGE:[^\]]+)?\]\]", lesson_text or ""))
    for index, match in enumerate(matches):
        page = match.group(1)
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(lesson_text or "")
        pages[page] = (lesson_text or "")[start:end]
    return pages


def page_number_from_url(image_url: str) -> str:
    match = re.search(r"page_(\d{3})\.png", str(image_url or ""))
    return str(int(match.group(1))) if match else ""


def crop_match_score(row: dict, page_text: str) -> int:
    question_terms = {term.lower() for term in keywords(str(row.get("question", "")), 12)}
    page_terms = {term.lower() for term in keywords(page_text, 30)}
    overlap = question_terms & page_terms
    score = len(overlap)
    if re.search(r"\b(?:diagram|graph|chart|figure|flow\s*chart|flowchart|map)\b", page_text, flags=re.IGNORECASE):
        score += 1
    return score


def attach_auto_crops(rows: list[dict], crop_index: dict[str, list[str]], lesson_text: str = "", force_visual: bool = False) -> list[dict]:
    page_texts = source_page_texts(lesson_text)
    for row in rows:
        q_type = str(row.get("question_type", "")).strip().lower()
        if force_visual and q_type in {"diagram", "graph"} and not row_needs_visual(row):
            question = str(row.get("question", "") or "").strip()
            visual_name = "graph" if q_type == "graph" else "diagram"
            row["question"] = f"Study the given {visual_name} and answer: {question}"
        review_reason = "" if force_visual and q_type in {"diagram", "graph"} else visual_review_reason(row)
        if review_reason:
            row["source_image_url"] = ""
            row["source_excerpt"] = review_reason
            row["review_status"] = "Needs Review"
            continue
        if not row_needs_visual(row):
            if str(row.get("question_type", "")).strip().lower() not in {"diagram", "graph"}:
                row["source_image_url"] = ""
                row["source_page"] = ""
            continue
        if str(row.get("source_image_url") or "").startswith("/textbook-crops/"):
            row["source_excerpt"] = row.get("source_excerpt") or "Approved crop used for this visual question."
            continue
        page_url = str(row.get("source_image_url") or "")
        page = str(row.get("source_page") or "").strip() or page_number_from_url(page_url)
        candidates: list[tuple[int, str, str]] = []
        for source_url, source_crops in crop_index.items():
            source_page = page_number_from_url(source_url)
            if page and source_page != page:
                continue
            score = crop_match_score(row, page_texts.get(source_page, ""))
            for crop_url in source_crops:
                candidates.append((score, source_page, crop_url))
        candidates.sort(reverse=True)
        if candidates and (candidates[0][0] >= 2 or page or (force_visual and q_type in {"diagram", "graph"})):
            score, matched_page, crop_url = candidates[0]
            row["source_page"] = matched_page or page
            row["source_image_url"] = crop_url
            if score >= 2:
                row["source_excerpt"] = f"Auto-attached image after keyword match score {score}. Verify crop before export."
            else:
                row["source_excerpt"] = "Needs image review: attached the closest crop from the selected source page."
                row["review_status"] = "Needs Review"
        else:
            row["source_image_url"] = ""
            row["source_excerpt"] = "Needs image review: no high-confidence matching crop found."
            row["review_status"] = "Needs Review"
    return rows


def approved_crop_index(crop_index: dict[str, list[str]], approved_urls: list[str] | None) -> dict[str, list[str]]:
    if not approved_urls:
        return crop_index
    approved = {str(url).strip() for url in approved_urls if str(url).strip()}
    if not approved:
        return {}
    filtered: dict[str, list[str]] = {}
    for source_url, crops in (crop_index or {}).items():
        kept = [crop for crop in crops if crop in approved]
        if kept:
            filtered[source_url] = kept
    return filtered


def visual_inputs_from_crops(crop_index: dict[str, list[str]], limit: int) -> list[dict]:
    inputs: list[dict] = []
    for source_url, crops in (crop_index or {}).items():
        source_page = page_number_from_url(source_url)
        for crop_url in crops:
            image_path = public_image_to_path(crop_url)
            if not image_path or not image_path.exists():
                continue
            inputs.append({
                "source_page": source_page,
                "source_image_url": crop_url,
                "image_path": str(image_path),
            })
            if len(inputs) >= limit:
                return inputs
    return inputs

def lesson_diagram_stats(text: str) -> dict:
    matches = DIAGRAM_TERMS_RE.findall(text or "")
    page_hits = set()
    current_page = ""
    for line in (text or "").splitlines():
        marker = re.match(r"\[\[PAGE:(\d+)", line.strip())
        if marker:
            current_page = marker.group(1)
            continue
        if current_page and DIAGRAM_TERMS_RE.search(line):
            page_hits.add(current_page)
    return {
        "diagram_available": bool(matches),
        "diagram_mentions": len(matches),
        "diagram_pages": sorted(page_hits, key=lambda value: int(value))[:12],
    }

def lesson_to_dict(lesson: LessonBlock, crop_index: dict[str, list[str]] | None = None) -> dict:
    diagram_stats = lesson_diagram_stats(lesson.text)
    crop_index = crop_index or {}
    crops = lesson_crop_urls(lesson.text, crop_index)
    source_pages = lesson_source_pages(lesson.text, crop_index)
    return {
        "lesson_no": lesson.lesson_no,
        "lesson_name": lesson.lesson_name,
        "word_count": len(lesson.text.split()),
        "auto_crop_count": len(crops),
        "sample_crop_url": crops[0] if crops else "",
        "source_pages": source_pages,
        **diagram_stats,
    }


def load_textbook_session(session_id: str) -> dict:
    session_path = TEXTBOOK_DIR / f"{session_id}.json"
    if not session_path.exists():
        raise HTTPException(status_code=404, detail="Textbook session not found. Analyze the textbook again.")
    return json.loads(session_path.read_text(encoding="utf-8"))


@app.post("/detect-lessons")
async def detect_textbook_lessons(
    files: list[UploadFile] = File(...),
    subject_name: str = Form("Textbook"),
    fallback_lesson_no: str = Form("Lesson01"),
    fallback_lesson_name: str = Form("Full textbook"),
):
    session_id = str(uuid.uuid4())
    extracted, errors = [], []
    for upload in files:
        try:
            extracted.append(await read_upload(upload, session_id))
        except Exception as exc:
            errors.append(f"{upload.filename}: {exc}")

    source_text = combine_texts(extracted)
    if not source_text.strip():
        raise HTTPException(status_code=400, detail=errors[0] if errors else "No readable textbook text found.")

    lessons = detect_lessons(source_text, fallback_lesson_no, fallback_lesson_name)
    crop_index = build_auto_crop_index(session_id)
    session_path = TEXTBOOK_DIR / f"{session_id}.json"
    session_path.write_text(
        json.dumps(
            {
                "session_id": session_id,
                "subject_name": subject_name,
                "lessons": [
                    {
                        "lesson_no": lesson.lesson_no,
                        "lesson_name": lesson.lesson_name,
                        "text": lesson.text,
                        "word_count": len(lesson.text.split()),
                    }
                    for lesson in lessons
                ],
                "auto_crops": crop_index,
                "file_errors": errors,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    return {
        "session_id": session_id,
        "lessons": [lesson_to_dict(lesson, crop_index) for lesson in lessons],
        "stats": {
            "files_read": len(extracted),
            "characters": len(source_text),
            "lessons": len(lessons),
            "file_errors": errors,
            "auto_crops": sum(len(crops) for crops in crop_index.values()),
        },
    }


@app.post("/generate-lesson")
async def generate_single_lesson(payload: dict = Body(...)):
    session_id = payload.get("session_id")
    lesson_no = payload.get("lesson_no")
    settings_payload = payload.get("settings") or {}
    if not session_id or not lesson_no:
        raise HTTPException(status_code=400, detail="session_id and lesson_no are required.")

    session = load_textbook_session(session_id)
    lesson_data = next((lesson for lesson in session.get("lessons", []) if lesson.get("lesson_no") == lesson_no), None)
    if not lesson_data:
        raise HTTPException(status_code=404, detail="Lesson not found in analyzed textbook.")

    requested_question_types = list(settings_payload.get("question_types") or ["objective", "very_short", "short", "long", "elaborative"])
    visual_source_enabled = bool(settings_payload.get("visual_source_enabled", False))
    approved_crop_urls = settings_payload.get("approved_crop_urls")
    if visual_source_enabled and approved_crop_urls and not any(q_type in {"diagram", "graph"} for q_type in requested_question_types):
        requested_question_types.append("diagram")

    total_count = max(1, min(int(settings_payload.get("count", 30)), 100))
    crop_index = approved_crop_index(session.get("auto_crops") or {}, approved_crop_urls)
    visual_types = [q_type for q_type in requested_question_types if q_type in {"diagram", "graph"}]
    normal_types = [q_type for q_type in requested_question_types if q_type not in {"diagram", "graph"}]
    visual_inputs = visual_inputs_from_crops(crop_index, total_count) if visual_source_enabled and visual_types else []
    if visual_source_enabled and visual_types and not visual_inputs:
        raise HTTPException(status_code=400, detail="Approve at least one visual crop before generating diagram/graph questions from images.")

    visual_count = min(total_count, len(visual_inputs)) if visual_inputs else 0
    normal_count = total_count - visual_count if normal_types else 0

    base_config = dict(
        difficulty=settings_payload.get("difficulty", "mixed"),
        subject_name=settings_payload.get("subject_name") or session.get("subject_name") or "Textbook",
        lesson_name=lesson_data.get("lesson_name") or lesson_no,
        lesson_no=lesson_no,
        bloom_tags=settings_payload.get("bloom_tags") or ["knowledge", "understanding", "application"],
        language=settings_payload.get("language", "English"),
        include_answers=bool(settings_payload.get("include_answers", True)),
        visual_source_enabled=visual_source_enabled,
    )
    config = GenerationSettings(
        count=total_count,
        question_types=requested_question_types,
        **base_config,
    )
    try:
        rows: list[dict] = []
        engines: set[str] = set()
        if normal_count:
            normal_config = GenerationSettings(
                count=normal_count,
                question_types=normal_types,
                **{**base_config, "visual_source_enabled": False},
            )
            normal_rows, normal_engine = await generate_questions(lesson_data.get("text", ""), normal_config)
            rows.extend(normal_rows)
            engines.add(normal_engine)
        if visual_count:
            visual_config = GenerationSettings(
                count=visual_count,
                question_types=visual_types,
                **base_config,
            )
            visual_rows, visual_engine = await generate_visual_questions_from_images(
                lesson_data.get("text", ""),
                visual_config,
                visual_inputs[:visual_count],
            )
            rows.extend(visual_rows)
            engines.add(visual_engine)
        if not rows:
            rows, engine = await generate_questions(lesson_data.get("text", ""), config)
        else:
            engine = "gemini" if engines == {"gemini"} else "mixed" if "gemini" in engines else (next(iter(engines)) if engines else "gemini")
    except AIUnavailableError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    for index, row in enumerate(rows, start=1):
        row["question_number"] = index
    clean_rows = [{column: row.get(column, "") for column in EXPORT_COLUMNS} for row in rows[:total_count]]
    clean_rows = attach_auto_crops(clean_rows, crop_index, lesson_data.get("text", ""), force_visual=False)
    return {
        "session_id": session_id,
        "engine": engine,
        "lesson": lesson_to_dict(
            LessonBlock(lesson_no=lesson_no, lesson_name=config.lesson_name, text=lesson_data.get("text", "")),
            session.get("auto_crops") or {},
        ),
        "rows": clean_rows,
        "stats": {
            "questions": len(clean_rows),
            "difficulty": config.difficulty,
        },
    }


@app.post("/regenerate-row")
async def regenerate_row(payload: dict = Body(...)):
    session_id = payload.get("session_id")
    lesson_no = payload.get("lesson_no")
    source_row = payload.get("row") or {}
    settings_payload = payload.get("settings") or {}
    if not session_id or not lesson_no:
        raise HTTPException(status_code=400, detail="session_id and lesson_no are required.")

    session = load_textbook_session(session_id)
    lesson_data = next((lesson for lesson in session.get("lessons", []) if lesson.get("lesson_no") == lesson_no), None)
    if not lesson_data:
        raise HTTPException(status_code=404, detail="Lesson not found in analyzed textbook.")

    q_type = source_row.get("question_type") or (settings_payload.get("question_types") or ["objective"])[0]
    q_type = str(q_type or "objective").strip().lower() or "objective"
    approved_crop_urls = settings_payload.get("approved_crop_urls")
    crop_index = approved_crop_index(session.get("auto_crops") or {}, approved_crop_urls)
    config = GenerationSettings(
        count=1,
        question_types=[q_type],
        difficulty=source_row.get("difficulty") or settings_payload.get("difficulty", "mixed"),
        subject_name=settings_payload.get("subject_name") or session.get("subject_name") or "Textbook",
        lesson_name=lesson_data.get("lesson_name") or lesson_no,
        lesson_no=lesson_no,
        bloom_tags=settings_payload.get("bloom_tags") or ["knowledge", "understanding", "application"],
        language=settings_payload.get("language", "English"),
        include_answers=bool(settings_payload.get("include_answers", True)),
        visual_source_enabled=bool(settings_payload.get("visual_source_enabled", False)),
    )
    try:
        if config.visual_source_enabled and q_type in {"diagram", "graph"}:
            source_image_url = str(source_row.get("source_image_url") or "").strip()
            visual_inputs = []
            if source_image_url.startswith("/textbook-crops/"):
                image_path = public_image_to_path(source_image_url)
                if image_path and image_path.exists():
                    visual_inputs.append({
                        "source_page": str(source_row.get("source_page") or page_number_from_url(source_image_url)),
                        "source_image_url": source_image_url,
                        "image_path": str(image_path),
                    })
            if not visual_inputs:
                visual_inputs = visual_inputs_from_crops(crop_index, 1)
            if not visual_inputs:
                raise HTTPException(status_code=400, detail="Approve at least one visual crop before regenerating this diagram/graph question.")
            rows, engine = await generate_visual_questions_from_images(lesson_data.get("text", ""), config, visual_inputs[:1])
        else:
            rows, engine = await generate_questions(lesson_data.get("text", ""), config)
    except AIUnavailableError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    row = rows[0] if rows else {}
    clean_row = {column: row.get(column, "") for column in EXPORT_COLUMNS}
    clean_row["question_number"] = source_row.get("question_number", clean_row.get("question_number", ""))
    clean_row = attach_auto_crops([clean_row], crop_index, lesson_data.get("text", ""), force_visual=config.visual_source_enabled)[0]
    return {"engine": engine, "row": clean_row}


@app.post("/generate-missing-answers")
async def generate_missing_answers(payload: dict = Body(...)):
    session_id = payload.get("session_id")
    lesson_no = payload.get("lesson_no")
    rows = payload.get("rows") or []
    settings_payload = payload.get("settings") or {}
    if not session_id or not lesson_no:
        raise HTTPException(status_code=400, detail="session_id and lesson_no are required.")
    if not isinstance(rows, list) or not rows:
        raise HTTPException(status_code=400, detail="No rows supplied for answer generation.")

    session = load_textbook_session(session_id)
    lesson_data = next((lesson for lesson in session.get("lessons", []) if lesson.get("lesson_no") == lesson_no), None)
    if not lesson_data:
        raise HTTPException(status_code=404, detail="Lesson not found in analyzed textbook.")

    missing_rows = []
    for row in rows:
        if not isinstance(row, dict) or str(row.get("AI answer") or "").strip():
            continue
        clean_row = {column: row.get(column, "") for column in EXPORT_COLUMNS}
        image_path = public_image_to_path(clean_row.get("source_image_url", ""))
        if image_path and image_path.exists():
            clean_row["__answer_image_path"] = str(image_path)
        missing_rows.append(clean_row)

    if not missing_rows:
        return {"engine": "none", "answers": []}

    config = GenerationSettings(
        count=len(missing_rows),
        question_types=list(settings_payload.get("question_types") or ["objective", "very_short", "short", "long", "elaborative"]),
        difficulty=settings_payload.get("difficulty", "mixed"),
        subject_name=settings_payload.get("subject_name") or session.get("subject_name") or "Textbook",
        lesson_name=lesson_data.get("lesson_name") or lesson_no,
        lesson_no=lesson_no,
        bloom_tags=settings_payload.get("bloom_tags") or ["knowledge", "understanding", "application"],
        language=settings_payload.get("language", "English"),
        include_answers=True,
        visual_source_enabled=bool(settings_payload.get("visual_source_enabled", False)),
    )
    try:
        answers, engine = await generate_answers_for_rows(missing_rows, lesson_data.get("text", ""), config)
    except AIUnavailableError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    return {"engine": engine, "answers": answers}


@app.post("/crop-source-image")
async def crop_source_image(payload: dict = Body(...)):
    image_url = str(payload.get("image_url") or "")
    crop = payload.get("crop") or {}
    if not image_url.startswith("/textbook-pages/"):
        raise HTTPException(status_code=400, detail="Only textbook page images can be cropped.")

    relative = image_url.replace("/textbook-pages/", "", 1).split("?", 1)[0].strip("/")
    image_path = (PAGE_IMAGE_DIR / relative).resolve()
    if PAGE_IMAGE_DIR.resolve() not in image_path.parents:
        raise HTTPException(status_code=400, detail="Invalid image path.")
    if not image_path.exists():
        raise HTTPException(status_code=404, detail="Source image not found.")

    try:
        from PIL import Image
    except Exception as exc:
        raise HTTPException(status_code=500, detail="Pillow is required for cropping.") from exc

    with Image.open(image_path) as image:
        width, height = image.size
        unit = crop.get("unit", "%")
        x = float(crop.get("x", 0))
        y = float(crop.get("y", 0))
        w = float(crop.get("width", 0))
        h = float(crop.get("height", 0))
        if unit == "%":
            left = int(width * x / 100)
            top = int(height * y / 100)
            right = int(width * (x + w) / 100)
            bottom = int(height * (y + h) / 100)
        else:
            left, top, right, bottom = int(x), int(y), int(x + w), int(y + h)
        left = max(0, min(left, width - 1))
        top = max(0, min(top, height - 1))
        right = max(left + 1, min(right, width))
        bottom = max(top + 1, min(bottom, height))
        cropped = image.crop((left, top, right, bottom))
        session_part = safe_path_part(str(payload.get("session_id") or "manual"))
        target_dir = CROP_IMAGE_DIR / session_part
        target_dir.mkdir(parents=True, exist_ok=True)
        filename = f"crop_{uuid.uuid4().hex[:10]}.png"
        cropped.save(target_dir / filename)
    return {"crop_url": f"/textbook-crops/{session_part}/{filename}"}


@app.post("/generate")
async def generate(files: list[UploadFile] = File(...), settings: str = Form(...)):
    try:
        payload = json.loads(settings)
        config = GenerationSettings(
            count=max(1, min(int(payload.get("count", 20)), 100)),
            question_types=payload.get("question_types") or ["objective"],
            difficulty=payload.get("difficulty", "average"),
            subject_name=payload.get("subject_name", "Textbook"),
            lesson_name=payload.get("lesson_name", "Full textbook"),
            lesson_no=payload.get("lesson_no", "All"),
            bloom_tags=payload.get("bloom_tags") or ["knowledge", "understanding", "application"],
            language=payload.get("language", "English"),
            include_answers=bool(payload.get("include_answers", True)),
        )
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Invalid settings: {exc}") from exc

    extracted, errors = [], []
    ad_hoc_session_id = str(uuid.uuid4())
    for upload in files:
        try:
            extracted.append(await read_upload(upload, ad_hoc_session_id))
        except Exception as exc:
            errors.append(f"{upload.filename}: {exc}")

    source_text = combine_texts(extracted)
    if not source_text.strip():
        raise HTTPException(status_code=400, detail=errors[0] if errors else "No readable textbook text found.")

    lessons = detect_lessons(source_text, config.lesson_no, config.lesson_name)
    try:
        workbook_rows, engine = await generate_lesson_questions(lessons, config)
    except AIUnavailableError as exc:
        raise HTTPException(status_code=503, detail=str(exc)) from exc
    generation_id = str(uuid.uuid4())
    path = EXPORT_DIR / f"generated_questions_{generation_id}.xlsx"
    write_workbook(path, workbook_rows)
    session_path = SESSION_DIR / f"{generation_id}.json"
    session_path.write_text(json.dumps({"workbook": workbook_rows, "stats": {"engine": engine}}, ensure_ascii=False, indent=2), encoding="utf-8")
    preview_rows = [row for rows in workbook_rows.values() for row in rows[:5]]
    return {
        "generation_id": generation_id,
        "engine": engine,
        "questions": preview_rows,
        "workbook": workbook_rows,
        "download_url": f"/download/{generation_id}",
        "stats": {
            "files_read": len(extracted),
            "characters": len(source_text),
            "questions": sum(len(rows) for rows in workbook_rows.values()),
            "lessons": len(lessons),
            "lesson_list": [
                {
                    "lesson_no": lesson.lesson_no,
                    "lesson_name": lesson.lesson_name,
                    "questions": len(workbook_rows.get(lesson.lesson_no, [])),
                }
                for lesson in lessons
            ],
            "file_errors": errors,
        },
    }


@app.post("/export-reviewed")
def export_reviewed(payload: dict = Body(...)):
    workbook_rows = payload.get("workbook") or {}
    if not isinstance(workbook_rows, dict) or not workbook_rows:
        raise HTTPException(status_code=400, detail="No reviewed workbook data supplied.")
    cleaned = clean_workbook_rows(workbook_rows)
    if not cleaned:
        raise HTTPException(status_code=400, detail="Workbook has no valid rows.")
    generation_id = str(uuid.uuid4())
    path = EXPORT_DIR / f"reviewed_questions_{generation_id}.xlsx"
    write_workbook(path, final_export_rows(cleaned), final_export=True)
    return {"generation_id": generation_id, "download_url": f"/download-reviewed/{generation_id}"}


@app.post("/review-session")
def save_review_session(payload: dict = Body(...)):
    workbook_rows = payload.get("workbook") or {}
    if not isinstance(workbook_rows, dict) or not workbook_rows:
        raise HTTPException(status_code=400, detail="No workbook data supplied.")
    session_id = payload.get("session_id") or str(uuid.uuid4())
    session_path = SESSION_DIR / f"{session_id}.json"
    session_path.write_text(
        json.dumps(
            {
                "session_id": session_id,
                "workbook": workbook_rows,
                "stats": payload.get("stats", {}),
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    return {"session_id": session_id, "saved": True}


@app.get("/review-session/{session_id}")
def load_review_session(session_id: str):
    session_path = SESSION_DIR / f"{session_id}.json"
    if not session_path.exists():
        raise HTTPException(status_code=404, detail="Review session not found")
    return json.loads(session_path.read_text(encoding="utf-8"))


@app.get("/download/{generation_id}")
def download(generation_id: str):
    path = EXPORT_DIR / f"generated_questions_{generation_id}.xlsx"
    if not path.exists():
        raise HTTPException(status_code=404, detail="Export not found")
    safe_id = generation_id[:8]
    return FileResponse(
        path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"generated_textbook_questions_{safe_id}.xlsx",
    )


@app.get("/download-reviewed/{generation_id}")
def download_reviewed(generation_id: str):
    path = EXPORT_DIR / f"reviewed_questions_{generation_id}.xlsx"
    if not path.exists():
        raise HTTPException(status_code=404, detail="Export not found")
    safe_id = generation_id[:8]
    return FileResponse(
        path,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"reviewed_textbook_questions_{safe_id}.xlsx",
    )

